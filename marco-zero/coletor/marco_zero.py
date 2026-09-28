"""CLI portátil para remontar uma coleta do painel e confrontá-la com um DuckDB."""
import argparse
import csv
import json
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path

from painel import PublicPanel, FIELDS, discover, filter_values, normalize, save_json, select_column
from validacao import validate_pair

COLUMNS = ['competencia', 'tipo_territorio', 'codigo_territorio', 'regiao', 'grupamento', *FIELDS]


def select_months(period, available):
    if not available:
        raise ValueError('Painel sem competências disponíveis.')
    if 'competencias' in period:
        selected = {int(m) for m in period['competencias']}
    else:
        start = int(period.get('inicio', min(available)))
        end = max(available) if period.get('fim', 'ultima_disponivel') == 'ultima_disponivel' else int(period['fim'])
        if start > end:
            raise ValueError('Início posterior ao fim do período.')
        month_numbers = period.get('meses_do_ano', list(range(1, 13)))
        if not month_numbers or any(not isinstance(m, int) or not 1 <= m <= 12 for m in month_numbers):
            raise ValueError('meses_do_ano deve conter números de 1 a 12.')
        selected = {y * 100 + m for y in range(start // 100, end // 100 + 1)
                    for m in month_numbers if start <= y * 100 + m <= end}
        selected |= {int(m) for m in period.get('adicionais', [])}
    if not selected or any(not 1 <= m % 100 <= 12 for m in selected):
        raise ValueError('Lista de competências vazia ou inválida.')
    missing = selected - set(available)
    if missing:
        raise ValueError('Competências ainda não disponíveis no painel: ' + str(sorted(missing)))
    return sorted(selected)


def resolve_territories(panel, entries):
    if not entries:
        raise ValueError('Configure ao menos um território.')
    catalogs, result, seen = {}, [], set()
    for entry in entries:
        kind = entry.get('tipo')
        if kind == 'municipio':
            uf = int(entry['uf'])
            if uf not in catalogs:
                catalogs[uf] = panel.catalog(uf)
            candidates = [(int(code), name) for code, name in catalogs[uf]
                          if (int(code) == int(entry['codigo']) if entry.get('codigo') is not None
                              else normalize(name) == normalize(entry['nome']))]
            if len(candidates) != 1:
                raise ValueError(f'Município não encontrado ou ambíguo: {entry}. Use o comando catalogo. Códigos municipais do painel têm 6 dígitos.')
            code, name = candidates[0]
            if entry.get('nome') and normalize(entry['nome']) != normalize(name):
                raise ValueError(f'Nome e código não coincidem: {entry} / {name}')
        elif kind == 'uf':
            code = int(entry['codigo'])
            found = panel.query('uf_' + str(code), [select_column('g', 'UF')], [filter_values('g', 'Código UF', [code])])
            if len(found) != 1 or found[0][0] is None:
                raise ValueError('UF inválida: ' + str(code))
            name = found[0][0]
            if entry.get('nome') and normalize(entry['nome']) != normalize(name):
                raise ValueError('Nome da UF incompatível: ' + str(found))
        else:
            raise ValueError('tipo deve ser municipio ou uf.')
        if (kind, code) in seen:
            raise ValueError('Território duplicado: ' + str(entry))
        seen.add((kind, code))
        result.append({'tipo': kind, 'codigo': code, 'nome': name})
    return result


def export(rows, directory):
    from openpyxl import Workbook, load_workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font
    from openpyxl.worksheet.table import Table, TableStyleInfo
    with (directory / 'dados.csv').open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS, delimiter=';')
        writer.writeheader()
        writer.writerows(rows)
    save_json(directory / 'dados.json', rows)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Dados'
    ws.append(COLUMNS)
    for record in rows:
        ws.append([record[k] for k in COLUMNS])
    ws.freeze_panes = 'F2'
    for letter, width in [('A', 15), ('B', 19), ('C', 20), ('D', 34), ('E', 24), ('F', 16), ('G', 18), ('H', 16), ('I', 17)]:
        ws.column_dimensions[letter].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row[5:]:
            cell.number_format = '#,##0;[Red]-#,##0;0'
        if row[4].value == 'TOTAL':
            for cell in row:
                cell.font = Font(bold=True)
    ws['I1'].comment = Comment('Estoque mensal oficial. Vazio preserva null da fonte; estoque não é saldo.', 'Fonte')
    ws['E1'].comment = Comment('TOTAL inclui os grupamentos. UF inclui seus municípios. Não somar agregados com suas parcelas.', 'Estrutura')
    table = Table(displayName='CagedMarcoZero', ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(table)
    path = directory / 'dados.xlsx'
    wb.save(path)
    check = load_workbook(path, read_only=True, data_only=True)
    expected = [tuple(COLUMNS)] + [tuple(r[k] for k in COLUMNS) for r in rows]
    if list(check.active.values) != expected:
        raise ValueError('Falha na verificação do Excel após exportação.')
    check.close()
    with (directory / 'dados.csv').open(encoding='utf-8-sig', newline='') as f:
        if list(csv.reader(f, delimiter=';')) != [[str(v) if v is not None else '' for v in r] for r in expected]:
            raise ValueError('Falha na verificação do CSV após exportação.')


def report(directory, rows, checks, session, comparison):
    observations = [c for c in checks if c['estoque_exige_observacao']]
    text = ['# Marco zero do Novo Caged', '',
            f'- Atualização do modelo: {session["refresh"]}.',
            f'- Linhas exportadas: {len(rows)}.',
            f'- Combinações de território e competência: {len(checks)}.',
            '- Saldo e totais de movimentações: conferidos por scripts.',
            '- Valores vazios preservados; saldo e estoque são medidas diferentes.',
            f'- Recortes com estoque vazio ou soma diferente do total: {len(observations)}.',
            '', '## Comparação local', '',
            json.dumps(comparison, ensure_ascii=False, indent=2) if comparison else
            'Não realizada: nenhum DuckDB foi informado. Os resultados oficiais foram validados internamente.',
            '', '## Limites', '',
            'A coleta usa o endpoint público do relatório, sem garantia de estabilidade como API do MTE. '
            'As evidências registram a versão consultada. Divergência entre fontes não identifica sozinha sua causa. '
            'Estoques ausentes não são preenchidos pelo residual do total. Diferenças de null e zero ficam explícitas.',
            '', '## Arquivos', '',
            '- dados.csv, dados.xlsx e dados.json: resultados em formato longo.',
            '- validacao.json: conferências por território/mês.',
            '- descoberta/: modelo, esquema e sessão pública.',
            '- evidencias/: requisições e respostas originais.',
            '- comparacao_duckdb.csv e validacao_duckdb.json: quando a comparação é solicitada.',
            '', '## Próximo passo para a IA', '',
            'Leia as validações e eventuais diferenças; separe fatos comprovados de hipóteses. '
            'Não altere números para forçar concordância entre o mart e o painel. '
            'Use PROMPT_PARA_AGENTE.md para investigar os meses divergentes e entregar a pasta ao usuário.']
    (directory / 'RELATORIO.md').write_text('\n'.join(text) + '\n', encoding='utf-8')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('acao', choices=['executar', 'catalogo', 'comparar'])
    parser.add_argument('--config', default='config.json')
    parser.add_argument('--duckdb', help='Sobrescreve comparacao.duckdb; aberto somente leitura.')
    parser.add_argument('--uf', type=int, default=28, help='UF para catalogo.')
    parser.add_argument('--pasta', help='Coleta existente, necessária para comparar.')
    args = parser.parse_args()
    config_path = Path(args.config).resolve()
    config = json.loads(config_path.read_text(encoding='utf-8-sig'))
    compare_config = config.get('comparacao', {})
    database = args.duckdb or compare_config.get('duckdb')
    if database:
        database = str((config_path.parent / Path(database).expanduser()).resolve())
    if args.acao == 'comparar':
        from warehouse import compare
        if not args.pasta or not database:
            raise ValueError('comparar requer --pasta e --duckdb (ou comparacao.duckdb).')
        directory = Path(args.pasta).resolve()
        rows = json.loads((directory / 'dados.json').read_text(encoding='utf-8'))
        print(json.dumps(compare(compare_config, database, rows, directory), ensure_ascii=False, indent=2))
        return
    stamp = datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ') + '_' + uuid.uuid4().hex[:6]
    directory = (config_path.parent / config.get('saida', 'saida') / ('marco-zero-' + stamp)).resolve()
    directory.mkdir(parents=True)
    save_json(directory / 'status.json', {'status': 'em_andamento'})
    print('Pasta:', directory, flush=True)
    try:
        print('Descobrindo modelo público...', flush=True)
        session = discover(config['painel_url'], directory / 'descoberta', config.get('navegador', 'chromium'))
        panel = PublicPanel(session, directory / 'evidencias')
        if args.acao == 'catalogo':
            save_json(directory / 'municipios.json', panel.catalog(args.uf))
            save_json(directory / 'status.json', {'status': 'catalogo_concluido'})
            print('Catálogo salvo:', directory / 'municipios.json')
            return
        territories = resolve_territories(panel, config['territorios'])
        available = panel.dates()
        months = select_months(config['periodo'], available)
        # ultima_competencia_disponivel: o que o painel já incorporava na coleta. É o
        # `retificacoes_ate` do marco zero (normalizar.py --coleta); sem ele, FOR/EXC
        # posteriores seriam ignorados ou contados duas vezes.
        save_json(directory / 'recorte.json', {'territorios': territories, 'competencias': months,
                                             'ultima_competencia_disponivel': max(available),
                                             'atualizacao_modelo': session['refresh'],
                                             'painel': session['painel_url'], 'coleta_utc': stamp})
        rows, checks = [], []
        for territory in territories:
            for month in months:
                groups = panel.metrics(territory, month, True)
                totals = panel.metrics(territory, month, False)
                checks.append(validate_pair(groups, totals, f'{territory["nome"]}/{month}'))
                for group, *values in [['TOTAL', *totals[0]], *groups]:
                    record = {'competencia': month, 'tipo_territorio': territory['tipo'],
                              'codigo_territorio': territory['codigo'], 'regiao': territory['nome'],
                              'grupamento': group if group is not None else 'Não Identificado'}
                    record.update(dict(zip(FIELDS, values)))
                    rows.append(record)
                print('Conferido:', territory['nome'], month, flush=True)
        keys = [(r['competencia'], r['tipo_territorio'], r['codigo_territorio'], r['grupamento']) for r in rows]
        if len(keys) != len(set(keys)):
            raise ValueError('Duplicidade no grão da exportação.')
        panel.check_refresh()
        rows.sort(key=lambda r: (r['competencia'], r['regiao'], r['grupamento'] != 'TOTAL', r['grupamento']))
        export(rows, directory)
        save_json(directory / 'validacao.json', checks)
        comparison = None
        if database:
            from warehouse import compare
            comparison = compare(compare_config, database, rows, directory)
        report(directory, rows, checks, session, comparison)
        save_json(directory / 'status.json', {'status': 'concluido', 'linhas': len(rows),
                                             'recortes_conferidos': len(checks),
                                             'comparacao_duckdb': comparison is not None})
        print('CONCLUÍDO:', directory, flush=True)
    except Exception as exc:
        save_json(directory / 'status.json', {'status': 'falhou', 'erro': str(exc)})
        raise


if __name__ == '__main__':
    try:
        main()
    except Exception as exc:
        print('ERRO:', exc, file=sys.stderr)
        sys.exit(1)
